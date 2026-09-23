"""
Lecture Voice-to-Notes Generator
Main Streamlit Application
"""

import streamlit as st
import os
from pathlib import Path
import tempfile
import ollama
import json
import re
import shutil
from datetime import datetime

# Page configuration
st.set_page_config(
    page_title="Lecture Notes AI",
    page_icon="🎓",
    layout="wide",
    initial_sidebar_state="expanded"
)

st.markdown(
    """
    <style>
    div[data-testid="column"]:has(.sticky-chat-marker) {
        position: sticky;
        top: 1rem;
        align-self: flex-start;
        max-height: calc(100vh - 2rem);
        overflow-y: auto;
        padding-bottom: 1rem;
    }

    div[data-testid="column"]:has(.sticky-chat-marker) textarea {
        max-height: 180px;
    }

    .ai-panel-shell {
        border: 1px solid rgba(250, 250, 250, 0.18);
        border-radius: 10px;
        padding: 0.75rem 0.75rem 1rem;
        background: rgba(255, 255, 255, 0.035);
    }

    .ai-panel-title {
        font-size: 1.1rem;
        font-weight: 700;
        margin-bottom: 0.35rem;
    }
    </style>
    """,
    unsafe_allow_html=True,
)

# Import modules
from src.audio_processor import AudioProcessor
from src.transcriber import Transcriber
from src.summarizer import NoteSummarizer
from src.quiz_generator import QuizGenerator
from src.flashcard_generator import FlashcardGenerator
from src.exam_generator import ExamGenerator, LocalSolutionGrader
from src.ai_provider import get_available_providers, get_provider_info, PROVIDER_CONFIGS
from src.json_utils import extract_json
from utils.file_handler import FileHandler
from utils.pdf_exporter import build_notes_pdf
import config

# Initialize session state
if 'transcription' not in st.session_state:
    st.session_state.transcription = None
if 'notes' not in st.session_state:
    st.session_state.notes = None
if 'quiz' not in st.session_state:
    st.session_state.quiz = None
if 'flashcards' not in st.session_state:
    st.session_state.flashcards = None
if 'quiz_data' not in st.session_state:
    st.session_state.quiz_data = None
if 'flashcard_data' not in st.session_state:
    st.session_state.flashcard_data = None
if 'exam_data' not in st.session_state:
    st.session_state.exam_data = None
if 'practice_tasks' not in st.session_state:
    st.session_state.practice_tasks = {"topic": "", "quick": [], "solution": [], "feedback": None}
if 'detected_language' not in st.session_state:
    st.session_state.detected_language = None
if 'flashcard_index' not in st.session_state:
    st.session_state.flashcard_index = 0
if 'show_flashcard_back' not in st.session_state:
    st.session_state.show_flashcard_back = False
if 'quiz_submitted' not in st.session_state:
    st.session_state.quiz_submitted = False
if 'exam_feedback' not in st.session_state:
    st.session_state.exam_feedback = None
if 'llm_provider' not in st.session_state:
    st.session_state.llm_provider = config.DEFAULT_PROVIDER
if 'llm_model' not in st.session_state:
    st.session_state.llm_model = config.DEFAULT_MODEL
if 'llm_api_key' not in st.session_state:
    st.session_state.llm_api_key = None
if 'chat_messages' not in st.session_state:
    st.session_state.chat_messages = []
if 'current_lecture_title' not in st.session_state:
    st.session_state.current_lecture_title = "Untitled Lecture"
if 'selected_folder' not in st.session_state:
    st.session_state.selected_folder = "General"
if 'source_file' not in st.session_state:
    st.session_state.source_file = None
if 'current_saved_lecture_path' not in st.session_state:
    st.session_state.current_saved_lecture_path = None
if 'current_lecture_dirty' not in st.session_state:
    st.session_state.current_lecture_dirty = False
if 'selected_text_for_ai' not in st.session_state:
    st.session_state.selected_text_for_ai = ''
if 'active_section' not in st.session_state:
    st.session_state.active_section = 'Upload'
if 'workspace_mode' not in st.session_state:
    st.session_state.workspace_mode = 'Конспекты'
if 'lecture_pending_delete' not in st.session_state:
    st.session_state.lecture_pending_delete = None
if 'ai_panel_open' not in st.session_state:
    st.session_state.ai_panel_open = True
if 'quiz_difficulty' not in st.session_state:
    st.session_state.quiz_difficulty = "Medium"
if 'flashcard_difficulty' not in st.session_state:
    st.session_state.flashcard_difficulty = "Medium"
if 'exam_difficulty' not in st.session_state:
    st.session_state.exam_difficulty = "Medium"
if 'task_difficulty' not in st.session_state:
    st.session_state.task_difficulty = "Medium"
if 'quick_tasks_checked' not in st.session_state:
    st.session_state.quick_tasks_checked = False

def main():
    """Main application function"""
    
    # Header
    st.title("🎓 Lecture Voice-to-Notes Generator")
    st.markdown("Transform your lecture recordings into structured notes, quizzes, and flashcards!")
    render_current_lecture_status()
    
    # Sidebar
    with st.sidebar:
        st.header("📁 Library")
        ensure_library()
        folders = list_folders()
        if st.session_state.selected_folder not in folders:
            st.session_state.selected_folder = folders[0]

        selected_folder = st.selectbox(
            "Folder",
            folders,
            index=folders.index(st.session_state.selected_folder),
        )
        st.session_state.selected_folder = selected_folder

        new_folder = st.text_input("New folder name")
        if st.button("Create Folder", use_container_width=True):
            if new_folder.strip():
                create_folder(new_folder.strip())
                st.session_state.selected_folder = new_folder.strip()
                st.rerun()

        saved_lectures = list_saved_lectures(selected_folder)
        if saved_lectures:
            st.caption("Saved lectures")
            for index, lecture in enumerate(saved_lectures):
                is_current = str(lecture["path"]) == st.session_state.get("current_saved_lecture_path")
                label = lecture["title"]
                if lecture.get("saved_at"):
                    label = f"{label} · {lecture['saved_at']}"
                if is_current:
                    label = f"▶ {label}"
                open_col, delete_col = st.columns([5, 1])
                with open_col:
                    if st.button(label, key=f"open_saved_{index}", use_container_width=True):
                        load_saved_lecture(lecture["path"])
                        st.rerun()
                with delete_col:
                    if st.button("🗑️", key=f"delete_saved_{index}", help="Delete lecture", use_container_width=True):
                        st.session_state.lecture_pending_delete = str(lecture["path"])
                        st.rerun()

            pending_delete = st.session_state.get("lecture_pending_delete")
            if pending_delete:
                pending_path = Path(pending_delete)
                pending_title = next(
                    (item["title"] for item in saved_lectures if str(item["path"]) == pending_delete),
                    pending_path.parent.name,
                )
                st.warning(f"Delete '{pending_title}'?")
                confirm_col, cancel_col = st.columns(2)
                with confirm_col:
                    if st.button("Delete", type="primary", use_container_width=True):
                        delete_saved_lecture(pending_path)
                        st.session_state.lecture_pending_delete = None
                        st.success("Deleted.")
                        st.rerun()
                with cancel_col:
                    if st.button("Cancel", use_container_width=True):
                        st.session_state.lecture_pending_delete = None
                        st.rerun()
        else:
            st.caption("No saved lectures in this folder yet.")

        st.divider()
        st.header("Lecture")
        lecture_title = st.text_input("Lecture title", value=st.session_state.current_lecture_title)
        st.session_state.current_lecture_title = lecture_title.strip() or "Untitled Lecture"
        if st.button("Create New Lecture", use_container_width=True):
            start_new_lecture()
            st.rerun()
        
        # Output preferences
        st.divider()
        st.header("Options")
        generate_summary = st.checkbox("Generate Summary", value=True)
        generate_quiz = st.checkbox("Generate Quiz", value=True)
        generate_flashcards = st.checkbox("Generate Flashcards", value=True)
        num_questions = 10
        num_flashcards = 15
        
        # Quiz settings
        if generate_quiz:
            num_questions = st.slider("Number of Quiz Questions", 5, 20, 10)
            st.session_state.quiz_difficulty = st.selectbox(
                "Quiz Difficulty",
                ["Easy", "Medium", "Hard", "Exam"],
                index=["Easy", "Medium", "Hard", "Exam"].index(st.session_state.quiz_difficulty),
            )
        
        # Flashcard settings
        if generate_flashcards:
            num_flashcards = st.slider("Number of Flashcards", 5, 30, 15)
            st.session_state.flashcard_difficulty = st.selectbox(
                "Flashcard Difficulty",
                ["Easy", "Medium", "Hard", "Exam"],
                index=["Easy", "Medium", "Hard", "Exam"].index(st.session_state.flashcard_difficulty),
            )

        st.session_state.exam_difficulty = st.selectbox(
            "Mini Exam Difficulty",
            ["Medium", "Hard", "Mixed", "Final Exam"],
            index=["Medium", "Hard", "Mixed", "Final Exam"].index(st.session_state.exam_difficulty),
        )
        
        transcription_provider = config.DEFAULT_TRANSCRIPTION_PROVIDER
        llm_provider = config.DEFAULT_PROVIDER
        model_choice = config.DEFAULT_MODEL
        api_key = None
        groq_api_key = None
    
    if st.session_state.ai_panel_open:
        main_col, chat_col = st.columns([3, 1], gap="large")
    else:
        main_col, chat_col = st.columns([12, 1], gap="large")
    with main_col:
        workspace_mode = st.radio(
            "Mode",
            ["Конспекты", "Задачи"],
            horizontal=True,
            label_visibility="collapsed",
        )
        st.session_state.workspace_mode = workspace_mode
        if workspace_mode == "Конспекты":
            section = st.radio(
                "Section",
                ["Upload", "Transcription", "Notes", "Quiz", "Flashcards", "Mini Exam"],
                horizontal=True,
                label_visibility="collapsed",
            )
        else:
            section = "Tasks"
        st.session_state.active_section = section
    
    # Upload
    if section == "Upload":
        with main_col:
            st.header("Add Lecture Material")
        
            col1, col2 = st.columns([2, 1])
        
            with col1:
                input_mode = st.radio(
                    "Input type",
                    ["Audio", "File or photo", "Text"],
                    horizontal=True,
                )
                if st.session_state.source_file and st.session_state.source_file.get("name") and not st.session_state.source_file.get("bytes"):
                    st.info(f"Opened saved lecture source file: {st.session_state.source_file['name']}")
                uploaded_file = st.file_uploader(
                    "Choose a file",
                    type=['mp3', 'wav', 'm4a', 'mp4', 'mpeg', 'mpga', 'webm', 'png', 'jpg', 'jpeg', 'webp', 'txt', 'md', 'pdf', 'docx', 'pptx'],
                    help="Supported: audio, images, text, PDF, Word, PowerPoint"
                )
                pasted_text = ""
                if input_mode == "Text":
                    pasted_text = st.text_area(
                        "Paste lecture text",
                        height=260,
                        placeholder="Paste your lecture text here..."
                    )
            
                if uploaded_file:
                    if not st.session_state.transcription and not st.session_state.notes:
                        st.session_state.current_lecture_title = Path(uploaded_file.name).stem
                    st.session_state.source_file = {
                        "name": uploaded_file.name,
                        "bytes": uploaded_file.getvalue(),
                    }
                    st.success(f"✅ File uploaded: {uploaded_file.name}")
                    st.write(f"📊 File size: {uploaded_file.size / (1024*1024):.2f} MB")
                
                    suffix = Path(uploaded_file.name).suffix.lower()
                    if suffix in ['.mp3', '.wav', '.m4a', '.mp4', '.mpeg', '.mpga', '.webm']:
                        st.audio(uploaded_file)
                    elif suffix in ['.png', '.jpg', '.jpeg', '.webp']:
                        st.image(uploaded_file)
                
                    # Process button
                    if st.button("🚀 Start Processing", type="primary", use_container_width=True):
                        suffix = Path(uploaded_file.name).suffix.lower()
                        is_audio = suffix in ['.mp3', '.wav', '.m4a', '.mp4', '.mpeg', '.mpga', '.webm']
                        # Validate API keys based on selected providers
                        if is_audio and transcription_provider == "openai" and not api_key:
                            st.error("❌ Please enter your OpenAI API key in the sidebar!")
                        elif is_audio and transcription_provider == "groq" and not groq_api_key:
                            st.error("❌ Please enter your Groq API key in the sidebar! Get it FREE at https://console.groq.com")
                        elif llm_provider == "openai" and not api_key:
                            st.error("❌ Please enter your OpenAI API key for text generation!")
                        elif llm_provider == "groq" and not groq_api_key:
                            st.error("❌ Please enter your Groq API key for text generation!")
                        else:
                            # Determine which API key to use for each provider
                            trans_key = groq_api_key if transcription_provider == "groq" else api_key
                            llm_key = groq_api_key if llm_provider == "groq" else api_key

                            if is_audio:
                                process_audio(
                                    uploaded_file,
                                    trans_key, llm_key,
                                    transcription_provider, llm_provider,
                                    model_choice,
                                    generate_summary, generate_quiz, generate_flashcards,
                                    num_questions if generate_quiz else 0,
                                    num_flashcards if generate_flashcards else 0
                                )
                            else:
                                process_file_input(
                                    uploaded_file,
                                    llm_key,
                                    llm_provider,
                                    model_choice,
                                    generate_summary, generate_quiz, generate_flashcards,
                                    num_questions if generate_quiz else 0,
                                    num_flashcards if generate_flashcards else 0
                                )
                elif input_mode == "Text":
                    text_title = st.text_input("Lecture title", value=st.session_state.current_lecture_title)
                    st.session_state.current_lecture_title = text_title.strip() or "Untitled Lecture"
                    if st.button("🚀 Start Processing", type="primary", use_container_width=True):
                        if llm_provider == "openai" and not api_key:
                            st.error("❌ Please enter your OpenAI API key for text generation!")
                        elif llm_provider == "groq" and not groq_api_key:
                            st.error("❌ Please enter your Groq API key for text generation!")
                        elif not pasted_text.strip():
                            st.error("❌ Paste lecture text first.")
                        else:
                            st.session_state.source_file = None
                            if not st.session_state.transcription and not st.session_state.notes:
                                st.session_state.current_lecture_title = text_title.strip() or "Untitled Lecture"
                            llm_key = groq_api_key if llm_provider == "groq" else api_key
                            process_lecture_text(
                                pasted_text.strip(),
                                llm_key,
                                llm_provider,
                                model_choice,
                                generate_summary, generate_quiz, generate_flashcards,
                                num_questions if generate_quiz else 0,
                                num_flashcards if generate_flashcards else 0,
                                source_language=None
                            )
        
            with col2:
                st.markdown("### 💡 Tips")
                st.markdown("""
                - Audio works best when speech is clear
                - Photos should be sharp and well lit
                - PDFs, Word files, PowerPoint files, and pasted text are supported
                - Output stays in the lecture language
                """)
            
                st.markdown("### ⏱️ Processing Time")
                st.markdown("""
                - **Audio/photo extraction** can take a few minutes locally
                - **Notes**: usually under a minute
                - **Quiz, flashcards, exam**: a few minutes locally
                """)
    
    # Transcription
    if section == "Transcription":
        with main_col:
            st.header("📝 Lecture Transcription")
        
            if st.session_state.transcription:
                st.success("✅ Transcription completed!")
            
                col1, col2 = st.columns([3, 1])
                with col1:
                    st.text_area(
                        "Full Transcription",
                        st.session_state.transcription,
                        height=400,
                        help="The complete transcription of your lecture"
                    )
            
                with col2:
                    if st.session_state.detected_language:
                        st.metric("Detected Language", language_label(st.session_state.detected_language))
                    st.metric("Word Count", len(st.session_state.transcription.split()))
                    st.metric("Character Count", len(st.session_state.transcription))
                
                    # Download button
                    if st.download_button(
                        label="📥 Download Transcript",
                        data=st.session_state.transcription,
                        file_name="transcript.txt",
                        mime="text/plain"
                    ):
                        st.success("Downloaded!")
                    if st.button("Ask AI about transcript", use_container_width=True):
                        set_ai_selection("Transcript", st.session_state.transcription)
                        st.rerun()
            else:
                st.info("👆 Add and process lecture material to see the extracted text here.")
    
    # Notes
    if section == "Notes":
        with main_col:
            st.header("📋 Structured Notes")
        
            if st.session_state.notes:
                st.success("✅ Notes generated!")
            
                st.markdown(st.session_state.notes)
            
                pdf_data = build_notes_pdf(
                    st.session_state.notes,
                    st.session_state.transcription or "",
                    language_label(st.session_state.detected_language) if st.session_state.detected_language else None
                )

                col1, col2, col3 = st.columns(3)
                with col1:
                    if st.download_button(
                        label="📥 Download as TXT",
                        data=st.session_state.notes,
                        file_name="lecture_notes.txt",
                        mime="text/plain"
                    ):
                        st.success("Downloaded!")
            
                with col2:
                    if st.download_button(
                        label="📥 Download as MD",
                        data=st.session_state.notes,
                        file_name="lecture_notes.md",
                        mime="text/markdown"
                    ):
                        st.success("Downloaded!")
                with col3:
                    if st.download_button(
                        label="📥 Download Beautiful PDF",
                        data=pdf_data,
                        file_name="lecture_notes.pdf",
                        mime="application/pdf"
                    ):
                        st.success("Downloaded!")
                if st.button("Ask AI about notes"):
                    set_ai_selection("Notes", st.session_state.notes)
                    st.rerun()
            else:
                st.info("👆 Add and process lecture material to see the notes here.")
    
    # Quiz
    if section == "Quiz":
        with main_col:
            st.header("❓ Practice Quiz")
        
            if st.session_state.quiz:
                st.success("✅ Quiz generated!")
                quiz_controls = st.columns([1, 1])
                with quiz_controls[0]:
                    st.session_state.quiz_difficulty = st.selectbox(
                        "Difficulty",
                        ["Easy", "Medium", "Hard", "Exam"],
                        index=["Easy", "Medium", "Hard", "Exam"].index(st.session_state.quiz_difficulty),
                        key="quiz_tab_difficulty",
                    )
                with quiz_controls[1]:
                    if st.button("Regenerate Quiz", type="primary", use_container_width=True):
                        regenerate_quiz(num_questions)
                        st.rerun()
                if st.session_state.quiz_data:
                    render_interactive_quiz(st.session_state.quiz_data)
                else:
                    st.markdown(st.session_state.quiz)
            
                if st.download_button(
                    label="📥 Download Quiz",
                    data=st.session_state.quiz,
                    file_name="lecture_quiz.txt",
                    mime="text/plain"
                ):
                    st.success("Downloaded!")
            else:
                st.info("👆 Add and process lecture material to see the quiz here.")
    
    # Flashcards
    if section == "Flashcards":
        with main_col:
            st.header("🗂️ Study Flashcards")
        
            if st.session_state.flashcards:
                st.success("✅ Flashcards generated!")
                flashcard_controls = st.columns([1, 1])
                with flashcard_controls[0]:
                    st.session_state.flashcard_difficulty = st.selectbox(
                        "Difficulty",
                        ["Easy", "Medium", "Hard", "Exam"],
                        index=["Easy", "Medium", "Hard", "Exam"].index(st.session_state.flashcard_difficulty),
                        key="flashcard_tab_difficulty",
                    )
                with flashcard_controls[1]:
                    if st.button("Regenerate Flashcards", type="primary", use_container_width=True):
                        regenerate_flashcards(num_flashcards)
                        st.rerun()
                if st.session_state.flashcard_data:
                    render_interactive_flashcards(st.session_state.flashcard_data)
                else:
                    st.markdown(st.session_state.flashcards)
            
                if st.download_button(
                    label="📥 Download Flashcards",
                    data=st.session_state.flashcards,
                    file_name="lecture_flashcards.txt",
                    mime="text/plain"
                ):
                    st.success("Downloaded!")
            else:
                st.info("👆 Add and process lecture material to see the flashcards here.")

    # Mini Exam
    if section == "Mini Exam":
        with main_col:
            st.header("🧪 Mini Exam")
            if st.session_state.exam_data:
                exam_controls = st.columns([1, 1])
                with exam_controls[0]:
                    st.session_state.exam_difficulty = st.selectbox(
                        "Difficulty",
                        ["Medium", "Hard", "Mixed", "Final Exam"],
                        index=["Medium", "Hard", "Mixed", "Final Exam"].index(st.session_state.exam_difficulty),
                        key="exam_tab_difficulty",
                    )
                with exam_controls[1]:
                    if st.button("Regenerate Mini Exam", type="primary", use_container_width=True):
                        regenerate_exam()
                        st.rerun()
                render_mini_exam(st.session_state.exam_data)
            else:
                st.info("👆 Add and process lecture material to generate a mini exam here.")

    # Tasks
    if section == "Tasks":
        with main_col:
            render_practice_tasks()

    with chat_col:
        st.markdown('<span class="sticky-chat-marker"></span>', unsafe_allow_html=True)
        if st.session_state.ai_panel_open:
            with st.container(border=True):
                close_col, title_col = st.columns([1, 5])
                with close_col:
                    if st.button("×", key="close_ai_panel", help="Close AI panel", use_container_width=True):
                        st.session_state.ai_panel_open = False
                        st.rerun()
                with title_col:
                    st.markdown('<div class="ai-panel-title">AI Tutor</div>', unsafe_allow_html=True)
                render_tutor_chat(compact=True)
        else:
            if st.button("AI", key="open_ai_panel", help="Open AI panel", use_container_width=True):
                st.session_state.ai_panel_open = True
                st.rerun()

def process_audio(file, trans_api_key, llm_api_key, trans_provider, llm_provider, 
                 model, gen_summary, gen_quiz, gen_flashcards, num_q, num_f):
    """Process the uploaded audio file"""
    
    try:
        # Step 1: Save uploaded file temporarily
        with st.spinner("💾 Saving audio file..."):
            file_handler = FileHandler()
            audio_path = file_handler.save_uploaded_file(file)
            st.success("✅ File saved!")
        
        # Step 2: Process audio (if needed)
        with st.spinner("🎵 Processing audio..."):
            processor = AudioProcessor()
            processed_path = processor.process(audio_path)
            st.success("✅ Audio processed!")
        
        # Step 3: Transcribe
        provider_name = trans_provider.upper() if trans_provider != "local" else "Faster-Whisper (Local)"
        with st.spinner(f"🎤 Transcribing audio with {provider_name} (this may take a few minutes)..."):
            transcriber = Transcriber(
                api_key=trans_api_key,
                provider=trans_provider
            )
            result = transcriber.transcribe_detailed(processed_path)
            transcription = result["text"]
            st.session_state.detected_language = result.get("language")
            if st.session_state.detected_language:
                st.success(f"✅ Transcription complete using {provider_name}! Language: {language_label(st.session_state.detected_language)}")
            else:
                st.success(f"✅ Transcription complete using {provider_name}!")

        process_lecture_text(
            transcription,
            llm_api_key,
            llm_provider,
            model,
            gen_summary,
            gen_quiz,
            gen_flashcards,
            num_q,
            num_f,
            source_language=st.session_state.detected_language,
            preserve_transcription=True
        )
        
        # Clean up temporary files
        file_handler.cleanup(audio_path)
        if processed_path != audio_path:
            file_handler.cleanup(processed_path)
        
    except Exception as e:
        st.error(f"❌ Error: {str(e)}")
        st.exception(e)


def process_file_input(file, llm_api_key, llm_provider, model, gen_summary, gen_quiz, gen_flashcards, num_q, num_f):
    """Extract lecture text from a non-audio file and process it."""
    suffix = Path(file.name).suffix.lower()
    with tempfile.NamedTemporaryFile(delete=False, suffix=suffix) as temp_file:
        temp_file.write(file.getvalue())
        temp_path = temp_file.name

    try:
        with st.spinner("📄 Extracting lecture text locally..."):
            if suffix in [".png", ".jpg", ".jpeg", ".webp"]:
                lecture_text = extract_text_from_image(temp_path)
            else:
                lecture_text = extract_text_from_file(temp_path, suffix)

        if not lecture_text.strip():
            st.error("❌ No readable lecture text was found in this file.")
            return

        process_lecture_text(
            lecture_text.strip(),
            llm_api_key,
            llm_provider,
            model,
            gen_summary,
            gen_quiz,
            gen_flashcards,
            num_q,
            num_f,
            source_language=None
        )
    finally:
        Path(temp_path).unlink(missing_ok=True)


def process_lecture_text(lecture_text, llm_api_key, llm_provider, model, gen_summary, gen_quiz, gen_flashcards, num_q, num_f, source_language=None, preserve_transcription=False):
    """Generate notes, quiz, flashcards, and exam from already extracted lecture text."""
    existing_text = st.session_state.transcription or ""
    if not source_language and st.session_state.get("detected_language"):
        source_language = st.session_state.detected_language
    existing_saved_path = st.session_state.get("current_saved_lecture_path")
    merging_into_existing = bool(existing_text.strip() and lecture_text.strip())
    if merging_into_existing:
        with st.spinner("Merging new material with the current lecture and removing repeats..."):
            lecture_text = merge_lecture_material(existing_text, lecture_text, llm_api_key, llm_provider, model, source_language)
            st.success("✅ New material merged into the current lecture.")

    reset_generated_state()
    st.session_state.transcription = lecture_text
    if merging_into_existing and existing_saved_path:
        st.session_state.current_saved_lecture_path = existing_saved_path
        st.session_state.current_lecture_dirty = True
    st.session_state.detected_language = source_language
    st.session_state.llm_provider = llm_provider
    st.session_state.llm_model = model or config.DEFAULT_MODEL
    st.session_state.llm_api_key = llm_api_key
    st.session_state.chat_messages = []
    generation_language = language_label(source_language) if source_language else "the same language as the lecture text"

    # Step 4: Generate notes
    if gen_summary:
        provider_name = llm_provider.upper()
        with st.spinner(f"📝 Generating structured notes with {provider_name}..."):
            summarizer = NoteSummarizer(
                api_key=llm_api_key,
                provider=llm_provider,
                model=model
            )
            notes = summarizer.generate_notes(
                f"Write the notes in {generation_language}. Do not translate to English unless the lecture is English.\n\n{lecture_text}"
            )
            st.session_state.notes = notes
            st.success(f"✅ Notes generated using {provider_name}!")

    # Step 5: Generate quiz
    if gen_quiz:
        provider_name = llm_provider.upper()
        with st.spinner(f"❓ Creating quiz questions with {provider_name}..."):
            quiz_gen = QuizGenerator(
                api_key=llm_api_key,
                provider=llm_provider,
                model=model
            )
            quiz_data = quiz_gen.generate_interactive_quiz(
                lecture_text,
                num_q,
                generation_language,
                st.session_state.quiz_difficulty.lower(),
            )
            quiz = format_quiz_for_download(quiz_data)
            st.session_state.quiz_data = quiz_data
            st.session_state.quiz = quiz
            st.session_state.quiz_submitted = False
            st.success(f"✅ Quiz created using {provider_name}!")

    # Step 6: Generate flashcards
    if gen_flashcards:
        provider_name = llm_provider.upper()
        with st.spinner(f"🗂️ Making flashcards with {provider_name}..."):
            flashcard_gen = FlashcardGenerator(
                api_key=llm_api_key,
                provider=llm_provider,
                model=model
            )
            flashcard_data = flashcard_gen.generate_interactive_flashcards(
                lecture_text,
                num_f,
                generation_language,
                st.session_state.flashcard_difficulty.lower(),
            )
            flashcards = format_flashcards_for_download(flashcard_data)
            st.session_state.flashcard_data = flashcard_data
            st.session_state.flashcards = flashcards
            st.session_state.flashcard_index = 0
            st.session_state.show_flashcard_back = False
            st.success(f"✅ Flashcards ready using {provider_name}!")

    # Step 7: Generate mini exam
    provider_name = llm_provider.upper()
    with st.spinner(f"🧪 Creating mini exam with {provider_name}..."):
        exam_gen = ExamGenerator(
            api_key=llm_api_key,
            provider=llm_provider,
            model=model
        )
        try:
            st.session_state.exam_data = exam_gen.generate_exam(
                lecture_text,
                generation_language,
                difficulty=st.session_state.exam_difficulty.lower(),
            )
            st.session_state.exam_feedback = None
            st.success(f"✅ Mini exam ready using {provider_name}!")
        except Exception as exam_error:
            st.session_state.exam_data = []
            st.warning(f"Mini exam could not be generated as structured tasks: {exam_error}")

    st.balloons()
    st.success("🎉 All done! Check the tabs above for your results.")
    save_current_lecture(st.session_state.selected_folder, st.session_state.current_lecture_title)
    st.success("Saved automatically.")


def merge_lecture_material(existing_text, new_text, llm_api_key, llm_provider, model, source_language=None):
    from src.ai_provider import AIProvider

    language = language_label(source_language) if source_language else "the lecture language"
    provider = AIProvider(
        provider_type=llm_provider or config.DEFAULT_PROVIDER,
        api_key=llm_api_key,
        model=model or config.DEFAULT_MODEL,
    )
    prompt = (
        f"Merge two lecture materials into one complete lecture transcript in {language}. "
        "Keep all unique facts, formulas, examples, definitions, slide content, and explanations. "
        "Remove duplicated repeated content. Do not summarize. Keep the original lecture language. "
        "Return only the merged lecture text.\n\n"
        f"CURRENT LECTURE MATERIAL:\n{existing_text}\n\n"
        f"NEW MATERIAL TO ADD:\n{new_text}"
    )
    return provider.generate_text(prompt, max_tokens=5000, temperature=0.15).strip()


def regenerate_quiz(num_questions: int):
    if not st.session_state.transcription:
        st.warning("Add lecture material first.")
        return
    language = language_label(st.session_state.detected_language) if st.session_state.detected_language else "the same language as the lecture text"
    quiz_gen = QuizGenerator(
        api_key=st.session_state.llm_api_key,
        provider=st.session_state.llm_provider or config.DEFAULT_PROVIDER,
        model=st.session_state.llm_model or config.DEFAULT_MODEL,
    )
    with st.spinner("Generating a new quiz locally..."):
        quiz_data = quiz_gen.generate_interactive_quiz(
            st.session_state.transcription,
            num_questions,
            language,
            st.session_state.quiz_difficulty.lower(),
        )
    st.session_state.quiz_data = quiz_data
    st.session_state.quiz = format_quiz_for_download(quiz_data)
    st.session_state.quiz_submitted = False
    save_current_lecture(st.session_state.selected_folder, st.session_state.current_lecture_title)


def regenerate_flashcards(num_flashcards: int):
    if not st.session_state.transcription:
        st.warning("Add lecture material first.")
        return
    language = language_label(st.session_state.detected_language) if st.session_state.detected_language else "the same language as the lecture text"
    flashcard_gen = FlashcardGenerator(
        api_key=st.session_state.llm_api_key,
        provider=st.session_state.llm_provider or config.DEFAULT_PROVIDER,
        model=st.session_state.llm_model or config.DEFAULT_MODEL,
    )
    with st.spinner("Generating new flashcards locally..."):
        flashcard_data = flashcard_gen.generate_interactive_flashcards(
            st.session_state.transcription,
            num_flashcards,
            language,
            st.session_state.flashcard_difficulty.lower(),
        )
    st.session_state.flashcard_data = flashcard_data
    st.session_state.flashcards = format_flashcards_for_download(flashcard_data)
    st.session_state.flashcard_index = 0
    st.session_state.show_flashcard_back = False
    save_current_lecture(st.session_state.selected_folder, st.session_state.current_lecture_title)


def regenerate_exam():
    if not st.session_state.transcription:
        st.warning("Add lecture material first.")
        return
    language = language_label(st.session_state.detected_language) if st.session_state.detected_language else "the same language as the lecture text"
    exam_gen = ExamGenerator(
        api_key=st.session_state.llm_api_key,
        provider=st.session_state.llm_provider or config.DEFAULT_PROVIDER,
        model=st.session_state.llm_model or config.DEFAULT_MODEL,
    )
    with st.spinner("Generating a new mini exam locally..."):
        st.session_state.exam_data = exam_gen.generate_exam(
            st.session_state.transcription,
            language,
            difficulty=st.session_state.exam_difficulty.lower(),
        )
    st.session_state.exam_feedback = None
    save_current_lecture(st.session_state.selected_folder, st.session_state.current_lecture_title)


def render_practice_tasks():
    st.header("🧮 Tasks")
    st.caption("Generate extra practice tasks by topic or keywords.")

    topic = st.text_input(
        "Topic or keywords",
        value=st.session_state.practice_tasks.get("topic", ""),
        placeholder="For example: derivatives, German cases, photosynthesis...",
    )
    st.session_state.practice_tasks["topic"] = topic
    controls = st.columns([1, 1, 1])
    with controls[0]:
        st.session_state.task_difficulty = st.selectbox(
            "Difficulty",
            ["Easy", "Medium", "Hard", "Exam"],
            index=["Easy", "Medium", "Hard", "Exam"].index(st.session_state.task_difficulty),
        )
    with controls[1]:
        quick_count = st.number_input("Quick-answer tasks", 1, 10, 4)
    with controls[2]:
        solution_count = st.number_input("Photo/file tasks", 1, 10, 3)

    if st.button("Generate Tasks", type="primary"):
        try:
            generate_practice_tasks(topic, int(quick_count), int(solution_count))
            st.rerun()
        except Exception as error:
            st.error(f"Could not generate tasks: {error}")

    quick_tasks = st.session_state.practice_tasks.get("quick", [])
    solution_tasks = st.session_state.practice_tasks.get("solution", [])

    if quick_tasks:
        st.divider()
        st.subheader("Quick Answer")
        if st.button("Regenerate Quick-Answer Tasks"):
            try:
                generate_quick_tasks(topic, int(quick_count))
                st.rerun()
            except Exception as error:
                st.error(f"Could not regenerate quick-answer tasks: {error}")
        render_quick_tasks(quick_tasks)

    if solution_tasks:
        st.divider()
        st.subheader("Upload Solution")
        if st.button("Regenerate Photo/File Tasks"):
            try:
                generate_solution_tasks(topic, int(solution_count))
                st.rerun()
            except Exception as error:
                st.error(f"Could not regenerate photo/file tasks: {error}")
        render_solution_tasks(solution_tasks)


def generate_practice_tasks(topic: str, quick_count: int, solution_count: int):
    if not topic.strip() and not st.session_state.transcription:
        st.warning("Enter a topic or open a lecture first.")
        return
    generate_quick_tasks(topic, quick_count)
    generate_solution_tasks(topic, solution_count)


def generate_quick_tasks(topic: str, count: int):
    st.session_state.practice_tasks["topic"] = topic
    st.session_state.practice_tasks["quick"] = generate_task_block(
        topic,
        count,
        "quick-answer tasks where the student types a short answer",
        ["title", "prompt", "matrices", "correct_answer", "explanation"],
    )
    st.session_state.practice_tasks["feedback"] = None
    save_current_lecture(st.session_state.selected_folder, st.session_state.current_lecture_title)


def generate_solution_tasks(topic: str, count: int):
    st.session_state.practice_tasks["topic"] = topic
    st.session_state.practice_tasks["solution"] = generate_task_block(
        topic,
        count,
        "worked-solution tasks where the student uploads a photo or file with full reasoning",
        ["title", "prompt", "matrices", "expected_answer", "rubric", "points"],
    )
    st.session_state.practice_tasks["feedback"] = None
    save_current_lecture(st.session_state.selected_folder, st.session_state.current_lecture_title)


def generate_task_block(topic: str, count: int, task_kind: str, fields: list) -> list:
    from src.ai_provider import AIProvider

    language = language_label(st.session_state.detected_language) if st.session_state.detected_language else "the language of the topic or lecture"
    field_list = ", ".join(fields)
    provider = AIProvider(
        provider_type=st.session_state.llm_provider or config.DEFAULT_PROVIDER,
        api_key=st.session_state.llm_api_key,
        model=st.session_state.llm_model or config.DEFAULT_MODEL,
    )
    prompt = (
        f"Create {count} {st.session_state.task_difficulty.lower()} {task_kind} in {language}. "
        f"Topic/keywords: {topic or 'use the current lecture context'}. "
        f"Return only a valid JSON array, starting with [ and ending with ]. Do not wrap it in an object. "
        f"Each item must include these fields: {field_list}. "
        "The prompt must contain the full task statement and all given values. "
        "If the task uses matrices, include them both inside prompt as JSON-like matrices and in matrices as arrays. "
        "Make tasks concrete and solvable. Create a fresh set, not duplicates."
    )
    lecture_context = st.session_state.transcription or st.session_state.notes or ""
    with st.spinner("Generating tasks locally..."):
        response = provider.chat_completion(
            messages=[
                {"role": "system", "content": prompt},
                {"role": "user", "content": lecture_context[:16000]},
            ],
            max_tokens=2400,
            temperature=0.55,
        )
    return normalize_generated_tasks(response, fields)


def normalize_generated_tasks(response: str, fields: list) -> list:
    try:
        parsed = extract_json(response)
    except Exception:
        return tasks_from_plain_text(response, fields)

    if isinstance(parsed, dict):
        for key in ("tasks", "questions", "items", "exercises", "problems"):
            if isinstance(parsed.get(key), list):
                parsed = parsed[key]
                break

    if isinstance(parsed, dict):
        parsed = [parsed]

    if not isinstance(parsed, list):
        return tasks_from_plain_text(response, fields)

    normalized = []
    for index, item in enumerate(parsed, 1):
        if isinstance(item, dict):
            task = {field: item.get(field, "") for field in fields}
            task["title"] = task.get("title") or item.get("name") or f"Task {index}"
            task["prompt"] = task.get("prompt") or item.get("question") or item.get("task") or item.get("problem") or ""
            task["matrices"] = task.get("matrices") or item.get("matrix") or item.get("given_matrices") or item.get("data") or []
            if "correct_answer" in fields:
                task["correct_answer"] = task.get("correct_answer") or item.get("answer") or item.get("solution") or ""
            if "expected_answer" in fields:
                task["expected_answer"] = task.get("expected_answer") or item.get("answer") or item.get("solution") or ""
            if "points" in fields:
                task["points"] = task.get("points") or 10
            normalized.append(task)
        else:
            normalized.append(task_from_text(str(item), index, fields))

    return [task for task in normalized if task.get("prompt")]


def tasks_from_plain_text(text: str, fields: list) -> list:
    chunks = re.split(r"\n\s*(?:\d+[\).\:-]|\-\s+|Task\s+\d+[:\).-]*)\s*", text)
    chunks = [chunk.strip() for chunk in chunks if chunk.strip()]
    if not chunks:
        chunks = [text.strip()]
    return [task_from_text(chunk, index, fields) for index, chunk in enumerate(chunks, 1)]


def task_from_text(text: str, index: int, fields: list) -> dict:
    task = {field: "" for field in fields}
    task["title"] = f"Task {index}"
    task["prompt"] = text
    if "correct_answer" in fields:
        task["correct_answer"] = ""
        task["explanation"] = "Generated from an unstructured local model response."
    if "expected_answer" in fields:
        task["expected_answer"] = ""
        task["rubric"] = "Check whether the solution is complete, correct, and clearly reasoned."
        task["points"] = 10
    return task


def repair_tasks_for_display(tasks: list, quick: bool) -> list:
    repaired = []
    for index, task in enumerate(tasks, 1):
        task = dict(task) if isinstance(task, dict) else task_from_text(str(task), index, ["title", "prompt"])
        task["title"] = str(task.get("title") or f"Task {index}")
        task["prompt"] = build_task_prompt(task, index, quick)
        if not quick:
            task["points"] = task.get("points") or 10
            task["rubric"] = task.get("rubric") or "Completeness, correctness, and clear reasoning."
        repaired.append(task)
    return repaired


def build_task_prompt(task: dict, index: int, quick: bool) -> str:
    prompt = str(task.get("prompt") or "").strip()
    if prompt:
        return prompt

    for key in ("question", "task", "problem", "description", "exercise", "statement", "aufgabe"):
        value = task.get(key)
        if value:
            return str(value).strip()

    title = str(task.get("title") or f"Task {index}").strip()
    matrices = normalize_matrices(task.get("matrices"))
    if matrices:
        if quick:
            return f"Solve this task: {title}. Use the given matrices and enter the final answer."
        return f"Solve this task and show your full reasoning: {title}. Use the given matrices."

    if quick:
        answer = str(task.get("correct_answer") or task.get("answer") or "").strip()
        if answer:
            return f"Solve this task: {title}."
        return f"Solve this task: {title}. Provide a short answer."

    expected = str(task.get("expected_answer") or task.get("answer") or "").strip()
    rubric = format_rubric(task.get("rubric"))
    details = expected or rubric
    if details:
        return f"Solve this task and show your full reasoning: {title}.\n\nExpected focus: {details}"
    return f"Solve this task and show your full reasoning: {title}."


def render_rich_task_text(text: str):
    text = str(text or "").strip()
    if not text:
        st.warning("Task text was missing, so it was rebuilt from the available task data.")
        return

    parts = split_text_and_matrices(text)
    if not parts:
        st.markdown(text)
        return

    for kind, value in parts:
        if kind == "text" and value.strip():
            st.markdown(value.strip())
        elif kind == "matrix":
            st.table(value)


def render_task_matrices(matrices):
    normalized = normalize_matrices(matrices)
    if not normalized:
        return
    st.markdown("Given matrices:")
    columns = st.columns(min(len(normalized), 3))
    for index, matrix in enumerate(normalized):
        with columns[index % len(columns)]:
            st.caption(f"Matrix {index + 1}")
            st.table(matrix)


def normalize_matrices(matrices):
    if not matrices:
        return []
    if isinstance(matrices, str):
        parsed = parse_matrix_literal(matrices)
        return [parsed] if parsed else []
    if is_matrix(matrices):
        return [matrices]
    if isinstance(matrices, list):
        normalized = []
        for item in matrices:
            if is_matrix(item):
                normalized.append(item)
            elif isinstance(item, str):
                parsed = parse_matrix_literal(item)
                if parsed:
                    normalized.append(parsed)
        return normalized
    if isinstance(matrices, dict):
        return normalize_matrices(list(matrices.values()))
    return []


def is_matrix(value) -> bool:
    return isinstance(value, list) and value and all(isinstance(row, list) for row in value)


def split_text_and_matrices(text: str) -> list:
    parts = []
    position = 0
    for match in re.finditer(r"\[\s*\[[^\]]+\](?:\s*,\s*\[[^\]]+\])+\s*\]", text):
        matrix_text = match.group(0)
        matrix = parse_matrix_literal(matrix_text)
        if not matrix:
            continue
        if match.start() > position:
            parts.append(("text", text[position:match.start()]))
        parts.append(("matrix", matrix))
        position = match.end()
    if position < len(text):
        parts.append(("text", text[position:]))
    return parts


def parse_matrix_literal(matrix_text: str):
    try:
        matrix = json.loads(matrix_text)
    except json.JSONDecodeError:
        return None
    if not isinstance(matrix, list) or not matrix or not all(isinstance(row, list) for row in matrix):
        return None
    return matrix


def format_rubric(rubric) -> str:
    if isinstance(rubric, list):
        return "; ".join(str(item) for item in rubric)
    if isinstance(rubric, dict):
        return "; ".join(f"{key}: {value}" for key, value in rubric.items())
    return str(rubric or "")


def render_quick_tasks(tasks: list):
    correct_count = 0
    checked = st.session_state.get("quick_tasks_checked", False)
    tasks = repair_tasks_for_display(tasks, quick=True)
    for index, task in enumerate(tasks):
        with st.container(border=True):
            st.markdown(f"#### {index + 1}. {task.get('title', 'Task')}")
            render_rich_task_text(task.get("prompt", ""))
            render_task_matrices(task.get("matrices"))
            st.text_input("Your answer", key=f"quick_task_answer_{index}")
            if checked:
                given = str(st.session_state.get(f"quick_task_answer_{index}", "")).strip().lower()
                correct = str(task.get("correct_answer", "")).strip().lower()
                is_correct = bool(given) and (given in correct or correct in given)
                correct_count += 1 if is_correct else 0
                st.markdown(f"**{'Correct' if is_correct else 'Needs review'}**")
                st.markdown("Correct answer:")
                render_rich_task_text(str(task.get("correct_answer", "")))
                st.caption(task.get("explanation", ""))
    if st.button("Check Quick Answers", type="primary"):
        st.session_state.quick_tasks_checked = True
        st.rerun()
    if checked:
        st.metric("Score", f"{correct_count}/{len(tasks)}")


def render_solution_tasks(tasks: list):
    tasks = repair_tasks_for_display(tasks, quick=False)
    for index, task in enumerate(tasks):
        with st.expander(f"{index + 1}. {task.get('title', 'Task')} - {task.get('points', '')} points", expanded=index == 0):
            render_rich_task_text(task.get("prompt", ""))
            render_task_matrices(task.get("matrices"))
            if task.get("rubric"):
                st.caption(f"Rubric: {format_rubric(task.get('rubric'))}")

    solution_text = st.text_area("Paste your worked solution", height=160, key="practice_solution_text")
    solution_file = st.file_uploader(
        "Or upload a photo/file with your solution",
        type=["png", "jpg", "jpeg", "webp", "txt", "md", "pdf", "docx"],
        key="practice_solution_file",
    )
    if st.button("Check Uploaded Solution", type="primary"):
        st.session_state.practice_tasks["feedback"] = grade_practice_solution(tasks, solution_text, solution_file)
        save_current_lecture(st.session_state.selected_folder, st.session_state.current_lecture_title)
        st.rerun()
    if st.session_state.practice_tasks.get("feedback"):
        st.subheader("Feedback")
        st.markdown(st.session_state.practice_tasks["feedback"])


def grade_practice_solution(tasks: list, solution_text: str, solution_file):
    grader = LocalSolutionGrader()
    language = language_label(st.session_state.detected_language) if st.session_state.detected_language else None
    if solution_file is not None:
        suffix = Path(solution_file.name).suffix.lower()
        with tempfile.NamedTemporaryFile(delete=False, suffix=suffix) as temp_file:
            temp_file.write(solution_file.getvalue())
            temp_path = temp_file.name
        try:
            if suffix in [".png", ".jpg", ".jpeg", ".webp"]:
                return grader.grade_image(tasks, temp_path, language)
            return grader.grade_text(tasks, extract_text_from_file(temp_path, suffix), language)
        finally:
            Path(temp_path).unlink(missing_ok=True)
    if solution_text.strip():
        return grader.grade_text(tasks, solution_text, language)
    return "Add text or upload a file first."


def reset_generated_state():
    st.session_state.notes = None
    st.session_state.quiz = None
    st.session_state.flashcards = None
    st.session_state.quiz_data = None
    st.session_state.flashcard_data = None
    st.session_state.exam_data = None
    st.session_state.exam_feedback = None
    st.session_state.practice_tasks = {"topic": "", "quick": [], "solution": [], "feedback": None}
    st.session_state.quick_tasks_checked = False
    st.session_state.quiz_submitted = False
    st.session_state.flashcard_index = 0
    st.session_state.show_flashcard_back = False
    st.session_state.chat_messages = []
    st.session_state.current_saved_lecture_path = None


def start_new_lecture():
    st.session_state.transcription = None
    st.session_state.notes = None
    st.session_state.quiz = None
    st.session_state.flashcards = None
    st.session_state.quiz_data = None
    st.session_state.flashcard_data = None
    st.session_state.exam_data = None
    st.session_state.practice_tasks = {"topic": "", "quick": [], "solution": [], "feedback": None}
    st.session_state.quick_tasks_checked = False
    st.session_state.detected_language = None
    st.session_state.flashcard_index = 0
    st.session_state.show_flashcard_back = False
    st.session_state.quiz_submitted = False
    st.session_state.exam_feedback = None
    st.session_state.chat_messages = []
    st.session_state.current_lecture_title = "Untitled Lecture"
    st.session_state.source_file = None
    st.session_state.current_saved_lecture_path = None
    st.session_state.current_lecture_dirty = False
    st.session_state.selected_text_for_ai = ""


def render_current_lecture_status():
    saved_path = st.session_state.get("current_saved_lecture_path")
    if saved_path:
        folder = Path(saved_path).parent.parent.name
        status = "Unsaved changes" if st.session_state.get("current_lecture_dirty") else "Saved"
        st.info(f"Current saved lecture: {st.session_state.current_lecture_title} · Folder: {folder} · {status}")
    elif st.session_state.transcription or st.session_state.notes:
        st.info(f"Current lecture: {st.session_state.current_lecture_title} · Not saved yet")


def ensure_library():
    library_root().mkdir(exist_ok=True)
    (library_root() / "General").mkdir(exist_ok=True)


def library_root() -> Path:
    return Path(__file__).parent / "saved_lectures"


def create_folder(name: str):
    (library_root() / safe_name(name)).mkdir(parents=True, exist_ok=True)


def list_folders() -> list:
    ensure_library()
    folders = [path.name for path in library_root().iterdir() if path.is_dir()]
    return sorted(folders) or ["General"]


def list_saved_lectures(folder: str) -> list:
    folder_path = library_root() / safe_name(folder)
    if not folder_path.exists():
        return []
    lectures = []
    json_paths = list(folder_path.glob("*.json")) + list(folder_path.glob("*/lecture.json"))
    for path in sorted(json_paths, key=lambda item: item.stat().st_mtime, reverse=True):
        try:
            data = json.loads(path.read_text())
            lectures.append({
                "title": data.get("title") or path.stem,
                "saved_at": data.get("saved_at", ""),
                "path": path,
            })
        except json.JSONDecodeError:
            continue
    return lectures


def save_current_lecture(folder: str, title: str):
    folder_path = library_root() / safe_name(folder)
    folder_path.mkdir(parents=True, exist_ok=True)
    existing_path = st.session_state.get("current_saved_lecture_path")
    lecture_json_path = Path(existing_path) if existing_path else None
    if lecture_json_path and lecture_json_path.exists() and library_root().resolve() in lecture_json_path.resolve().parents:
        lecture_dir = lecture_json_path.parent
    else:
        timestamp = datetime.now().strftime("%Y%m%d-%H%M%S")
        lecture_dir = folder_path / f"{timestamp}-{safe_name(title)}"
        lecture_json_path = lecture_dir / "lecture.json"

    source_dir = lecture_dir / "source"
    source_dir.mkdir(parents=True, exist_ok=True)
    source_file_path = None
    previous_data = {}
    if lecture_json_path.exists():
        try:
            previous_data = json.loads(lecture_json_path.read_text())
        except json.JSONDecodeError:
            previous_data = {}
    source_file = st.session_state.get("source_file")
    if source_file and source_file.get("bytes"):
        source_file_path = source_dir / safe_filename(source_file.get("name") or "source")
        source_file_path.write_bytes(source_file["bytes"])
    elif previous_data.get("source_file"):
        source_file_path = lecture_dir / previous_data["source_file"]
    data = {
        "title": title,
        "saved_at": datetime.now().isoformat(timespec="seconds"),
        "language": st.session_state.detected_language,
        "source_file": str(source_file_path.relative_to(lecture_dir)) if source_file_path and source_file_path.exists() else previous_data.get("source_file"),
        "transcription": st.session_state.transcription,
        "notes": st.session_state.notes,
        "quiz": st.session_state.quiz,
        "flashcards": st.session_state.flashcards,
        "quiz_data": st.session_state.quiz_data,
        "flashcard_data": st.session_state.flashcard_data,
        "exam_data": st.session_state.exam_data,
        "exam_feedback": st.session_state.exam_feedback,
        "practice_tasks": st.session_state.practice_tasks,
        "chat_messages": st.session_state.chat_messages,
    }
    lecture_json_path.write_text(json.dumps(data, ensure_ascii=False, indent=2))
    st.session_state.current_saved_lecture_path = str(lecture_json_path)
    st.session_state.current_lecture_dirty = False


def load_saved_lecture(path: Path):
    data = json.loads(path.read_text())
    st.session_state.current_lecture_title = data.get("title") or path.stem
    st.session_state.detected_language = data.get("language")
    st.session_state.transcription = data.get("transcription")
    st.session_state.notes = data.get("notes")
    st.session_state.quiz = data.get("quiz")
    st.session_state.flashcards = data.get("flashcards")
    st.session_state.quiz_data = data.get("quiz_data")
    st.session_state.flashcard_data = data.get("flashcard_data")
    st.session_state.exam_data = data.get("exam_data")
    st.session_state.exam_feedback = data.get("exam_feedback")
    st.session_state.practice_tasks = data.get("practice_tasks") or {"topic": "", "quick": [], "solution": [], "feedback": None}
    st.session_state.chat_messages = data.get("chat_messages") or []
    st.session_state.current_saved_lecture_path = str(path)
    st.session_state.current_lecture_dirty = False
    source = data.get("source_file")
    st.session_state.source_file = {"name": Path(source).name, "bytes": None} if source else None
    st.session_state.quiz_submitted = False
    st.session_state.flashcard_index = 0
    st.session_state.show_flashcard_back = False


def delete_saved_lecture(path: Path):
    library = library_root().resolve()
    target = path.resolve()
    if library not in target.parents:
        raise ValueError("This file is outside the lecture library.")
    if target.name == "lecture.json" and target.parent != library:
        shutil.rmtree(target.parent)
    elif target.suffix == ".json":
        target.unlink(missing_ok=True)
    if st.session_state.get("current_saved_lecture_path") == str(path):
        st.session_state.current_saved_lecture_path = None
        st.session_state.current_lecture_dirty = False


def safe_name(name: str) -> str:
    cleaned = re.sub(r"[^A-Za-z0-9А-Яа-яЁё._ -]+", "", name).strip()
    cleaned = re.sub(r"\s+", " ", cleaned)
    return cleaned[:80] or "Untitled"


def safe_filename(name: str) -> str:
    path = Path(name)
    stem = safe_name(path.stem)
    suffix = re.sub(r"[^A-Za-z0-9.]+", "", path.suffix)[:12]
    return f"{stem}{suffix}"

def language_label(code: str) -> str:
    labels = {
        "de": "German",
        "en": "English",
        "ru": "Russian",
        "fr": "French",
        "es": "Spanish",
        "it": "Italian",
        "pt": "Portuguese",
        "nl": "Dutch",
        "pl": "Polish",
        "uk": "Ukrainian",
    }
    return labels.get(code, code)


def render_interactive_flashcards(cards: list):
    if not cards:
        st.info("No flashcards available.")
        return

    index = st.session_state.flashcard_index % len(cards)
    card = cards[index]
    st.caption(f"Card {index + 1} of {len(cards)}")
    st.subheader(card.get("topic", "Flashcard"))
    st.markdown(f"### {card.get('front', '')}")

    if st.session_state.show_flashcard_back:
        st.divider()
        st.markdown(card.get("back", ""))

    card_context = (
        f"Topic: {card.get('topic', 'Flashcard')}\n"
        f"Front: {card.get('front', '')}\n"
        f"Back: {card.get('back', '')}"
    )
    if st.button("Ask AI about this card", use_container_width=True):
        set_ai_selection("Flashcard", card_context)
        st.rerun()

    col1, col2, col3 = st.columns(3)
    with col1:
        if st.button("Previous", use_container_width=True):
            st.session_state.flashcard_index = (index - 1) % len(cards)
            st.session_state.show_flashcard_back = False
            st.rerun()
    with col2:
        if st.button("Flip", use_container_width=True):
            st.session_state.show_flashcard_back = not st.session_state.show_flashcard_back
            st.rerun()
    with col3:
        if st.button("Next", use_container_width=True):
            st.session_state.flashcard_index = (index + 1) % len(cards)
            st.session_state.show_flashcard_back = False
            st.rerun()


def render_interactive_quiz(questions: list):
    if not questions:
        st.info("No quiz questions available.")
        return

    answers = {}
    for i, question in enumerate(questions):
        st.markdown(f"#### {i + 1}. {question.get('question', '')}")
        key = f"quiz_answer_{i}"
        if question.get("type") == "multiple_choice":
            options = question.get("options") or []
            answers[i] = st.radio("Choose an answer", options, key=key, label_visibility="collapsed")
        else:
            answers[i] = st.text_area("Your answer", key=key)
        question_context = format_quiz_question_for_context(question, i + 1)
        if st.button("Ask AI about this question", key=f"ask_quiz_{i}"):
            set_ai_selection("Quiz question", question_context)
            st.rerun()

    if st.button("Submit Quiz", type="primary"):
        st.session_state.quiz_submitted = True

    if st.session_state.quiz_submitted:
        score = 0
        st.divider()
        st.subheader("Results")
        for i, question in enumerate(questions):
            correct = str(question.get("correct_answer", "")).strip()
            given = str(st.session_state.get(f"quiz_answer_{i}", "")).strip()
            is_correct = False
            if question.get("type") == "multiple_choice":
                is_correct = given == correct
            else:
                is_correct = bool(given) and (given.lower() in correct.lower() or correct.lower() in given.lower())
            score += 1 if is_correct else 0
            st.markdown(f"**{i + 1}. {'Correct' if is_correct else 'Needs review'}**")
            st.write(f"Your answer: {given or '-'}")
            st.write(f"Correct answer: {correct}")
            st.caption(question.get("explanation", ""))
        st.metric("Score", f"{score}/{len(questions)}")


def render_mini_exam(tasks: list):
    total_points = sum(int(task.get("points", 0)) for task in tasks)
    st.caption(f"Total: {total_points} points")
    for i, task in enumerate(tasks):
        with st.expander(f"{i + 1}. {task.get('title', 'Task')} - {task.get('difficulty', '')} - {task.get('points', 0)} points", expanded=i == 0):
            st.write(task.get("prompt", ""))
            task_context = (
                f"{i + 1}. {task.get('title', 'Task')}\n"
                f"Difficulty: {task.get('difficulty', '')}\n"
                f"Points: {task.get('points', 0)}\n"
                f"Task: {task.get('prompt', '')}\n"
                f"Expected answer: {task.get('expected_answer', '')}"
            )
            if st.button("Ask AI about this task", key=f"ask_exam_{i}"):
                set_ai_selection("Mini exam task", task_context)
                st.rerun()

    solution_text = st.text_area("Paste your solution here", height=180)
    solution_file = st.file_uploader(
        "Or upload a photo/file with your solution",
        type=["png", "jpg", "jpeg", "webp", "txt", "md", "pdf", "docx"],
    )

    if st.button("Check Solution", type="primary"):
        grader = LocalSolutionGrader()
        language = language_label(st.session_state.detected_language) if st.session_state.detected_language else None
        with st.spinner("Checking solution locally..."):
            if solution_file is not None:
                suffix = Path(solution_file.name).suffix.lower()
                with tempfile.NamedTemporaryFile(delete=False, suffix=suffix) as temp_file:
                    temp_file.write(solution_file.getvalue())
                    temp_path = temp_file.name
                try:
                    if suffix in [".png", ".jpg", ".jpeg", ".webp"]:
                        st.session_state.exam_feedback = grader.grade_image(tasks, temp_path, language)
                    else:
                        st.session_state.exam_feedback = grader.grade_text(tasks, extract_text_from_file(temp_path, suffix), language)
                finally:
                    Path(temp_path).unlink(missing_ok=True)
            elif solution_text.strip():
                st.session_state.exam_feedback = grader.grade_text(tasks, solution_text, language)
            else:
                st.warning("Add text or upload a file first.")

    if st.session_state.exam_feedback:
        st.divider()
        st.subheader("Feedback")
        st.markdown(st.session_state.exam_feedback)


def render_tutor_chat(compact: bool = False):
    st.caption(f"Context: {st.session_state.active_section}")
    if not st.session_state.transcription:
        st.info("Process a lecture first, then ask questions here.")
        return

    selected_fragment = st.text_area(
        "Selected fragment",
        value=st.session_state.selected_text_for_ai,
        height=100,
        placeholder="Paste a word, line, paragraph, quiz question, flashcard, or exam task here if you want to ask about a specific part."
    )
    st.session_state.selected_text_for_ai = selected_fragment

    for message in st.session_state.chat_messages:
        with st.chat_message(message["role"]):
            st.markdown(message["content"])

    chat_mode = st.radio(
        "Mode",
        ["Ask", "Edit current section"],
        horizontal=False,
        label_visibility="collapsed",
    )
    question = st.text_area(
        "Ask or tell AI what to change",
        key="tutor_question",
        height=90 if compact else 120,
        placeholder="Ask a question, or choose Edit and write what to change..."
    )
    send_clicked = st.button("Send", type="primary", use_container_width=True)

    if send_clicked and question.strip():
        question = question.strip()
        should_refresh_after_edit = chat_mode == "Edit current section"
        user_text = question
        if selected_fragment.strip():
            user_text = f"Selected fragment:\n{selected_fragment.strip()}\n\nQuestion:\n{question}"
        st.session_state.chat_messages.append({"role": "user", "content": user_text})
        with st.chat_message("user"):
            st.markdown(user_text)

        with st.chat_message("assistant"):
            with st.spinner("Thinking locally..."):
                try:
                    if chat_mode == "Edit current section":
                        answer = edit_current_section(user_text)
                    else:
                        answer = answer_tutor_question(user_text)
                except Exception as error:
                    answer = f"I could not apply that request: {error}"
                st.markdown(answer)
        st.session_state.chat_messages.append({"role": "assistant", "content": answer})
        if should_refresh_after_edit and answer.startswith("Done."):
            st.rerun()

    if st.button("Clear Chat", use_container_width=True):
        st.session_state.chat_messages = []
        st.session_state.selected_text_for_ai = ""
        st.rerun()


def answer_tutor_question(question: str) -> str:
    from src.ai_provider import AIProvider

    context = build_tutor_context()
    language = language_label(st.session_state.detected_language) if st.session_state.detected_language else "the language of the user's question"
    active_section = st.session_state.get("active_section", "Unknown")
    selected_text = st.session_state.get("selected_text_for_ai", "").strip()
    provider = AIProvider(
        provider_type=st.session_state.llm_provider or config.DEFAULT_PROVIDER,
        api_key=st.session_state.llm_api_key,
        model=st.session_state.llm_model or config.DEFAULT_MODEL,
    )
    messages = [
        {
            "role": "system",
            "content": (
                f"You are a patient study tutor. Answer in {language}. "
                f"The student is currently viewing the {active_section} section. "
                "Use the lecture context below and prioritize the currently selected fragment when it is provided. "
                "Explain step by step when useful. "
                "If the student asks about a quiz, flashcard, or exam task, explain the reasoning and not just the answer.\n\n"
                f"CURRENTLY SELECTED FRAGMENT:\n{selected_text or 'None'}\n\n"
                f"{context}"
            ),
        }
    ]
    messages.extend(st.session_state.chat_messages[-8:-1])
    messages.append({"role": "user", "content": question})
    return provider.chat_completion(messages, max_tokens=1400, temperature=0.4)


def edit_current_section(instruction: str) -> str:
    from src.ai_provider import AIProvider

    section = st.session_state.get("active_section", "Notes")
    language = language_label(st.session_state.detected_language) if st.session_state.detected_language else "the lecture language"
    selected_text = st.session_state.get("selected_text_for_ai", "").strip()
    provider = AIProvider(
        provider_type=st.session_state.llm_provider or config.DEFAULT_PROVIDER,
        api_key=st.session_state.llm_api_key,
        model=st.session_state.llm_model or config.DEFAULT_MODEL,
    )

    if section == "Notes":
        if not st.session_state.notes:
            return "There are no notes to edit yet."
        prompt = (
            f"Rewrite the lecture notes in {language} according to the student's instruction. "
            "Keep clean Markdown formatting, useful headings, emphasis, examples, formulas, and tables when helpful. "
            "Return only the updated notes, with no explanation before or after.\n\n"
            f"Student instruction:\n{instruction}\n\n"
            f"Selected fragment, if relevant:\n{selected_text or 'None'}\n\n"
            f"Current notes:\n{st.session_state.notes}\n\n"
            f"Lecture context:\n{st.session_state.transcription or ''}"
        )
        updated_notes = provider.generate_text(prompt, max_tokens=2600, temperature=0.35).strip()
        st.session_state.notes = updated_notes
        save_current_lecture(st.session_state.selected_folder, st.session_state.current_lecture_title)
        return "Done. I updated the notes."

    if section == "Quiz":
        if not st.session_state.quiz_data:
            return "There is no interactive quiz to edit yet."
        prompt = (
            f"Edit this interactive quiz in {language} according to the student's instruction. "
            "Return only valid JSON array. Each item must have: type, question, options, correct_answer, explanation. "
            "Use type 'multiple_choice' when options exist, otherwise 'short_answer'. "
            "Do not include Markdown fences or commentary.\n\n"
            f"Student instruction:\n{instruction}\n\n"
            f"Selected fragment, if relevant:\n{selected_text or 'None'}\n\n"
            f"Current quiz JSON:\n{json.dumps(st.session_state.quiz_data, ensure_ascii=False)}\n\n"
            f"Lecture context:\n{st.session_state.transcription or ''}"
        )
        response = provider.generate_text(prompt, max_tokens=2600, temperature=0.25)
        updated_quiz = extract_json(response)
        if not isinstance(updated_quiz, list):
            raise ValueError("AI returned quiz in the wrong format.")
        st.session_state.quiz_data = updated_quiz
        st.session_state.quiz = format_quiz_for_download(updated_quiz)
        st.session_state.quiz_submitted = False
        save_current_lecture(st.session_state.selected_folder, st.session_state.current_lecture_title)
        return "Done. I updated the quiz."

    if section == "Flashcards":
        if not st.session_state.flashcard_data:
            return "There are no interactive flashcards to edit yet."
        prompt = (
            f"Edit these flashcards in {language} according to the student's instruction. "
            "Return only valid JSON array. Each item must have: topic, front, back. "
            "Do not include Markdown fences or commentary.\n\n"
            f"Student instruction:\n{instruction}\n\n"
            f"Selected fragment, if relevant:\n{selected_text or 'None'}\n\n"
            f"Current flashcards JSON:\n{json.dumps(st.session_state.flashcard_data, ensure_ascii=False)}\n\n"
            f"Lecture context:\n{st.session_state.transcription or ''}"
        )
        response = provider.generate_text(prompt, max_tokens=2600, temperature=0.25)
        updated_cards = extract_json(response)
        if not isinstance(updated_cards, list):
            raise ValueError("AI returned flashcards in the wrong format.")
        st.session_state.flashcard_data = updated_cards
        st.session_state.flashcards = format_flashcards_for_download(updated_cards)
        st.session_state.flashcard_index = min(st.session_state.flashcard_index, max(len(updated_cards) - 1, 0))
        st.session_state.show_flashcard_back = False
        save_current_lecture(st.session_state.selected_folder, st.session_state.current_lecture_title)
        return "Done. I updated the flashcards."

    if section == "Mini Exam":
        if not st.session_state.exam_data:
            return "There is no mini exam to edit yet."
        prompt = (
            f"Edit this mini exam in {language} according to the student's instruction. "
            "Return only valid JSON array. Each item must have: title, difficulty, points, prompt, expected_answer. "
            "Do not include Markdown fences or commentary.\n\n"
            f"Student instruction:\n{instruction}\n\n"
            f"Selected fragment, if relevant:\n{selected_text or 'None'}\n\n"
            f"Current mini exam JSON:\n{json.dumps(st.session_state.exam_data, ensure_ascii=False)}\n\n"
            f"Lecture context:\n{st.session_state.transcription or ''}"
        )
        response = provider.generate_text(prompt, max_tokens=2600, temperature=0.25)
        updated_exam = extract_json(response)
        if not isinstance(updated_exam, list):
            raise ValueError("AI returned mini exam in the wrong format.")
        st.session_state.exam_data = updated_exam
        st.session_state.exam_feedback = None
        save_current_lecture(st.session_state.selected_folder, st.session_state.current_lecture_title)
        return "Done. I updated the mini exam."

    return "Open Notes, Quiz, Flashcards, or Mini Exam, then choose Edit current section."


def build_tutor_context() -> str:
    parts = [
        "LECTURE TEXT:",
        st.session_state.transcription or "",
        "\nNOTES:",
        st.session_state.notes or "",
        "\nQUIZ:",
        st.session_state.quiz or "",
        "\nFLASHCARDS:",
        st.session_state.flashcards or "",
        "\nMINI EXAM:",
        format_exam_for_context(st.session_state.exam_data or []),
        "\nTASKS:",
        json.dumps(st.session_state.practice_tasks or {}, ensure_ascii=False, indent=2),
    ]
    return "\n".join(parts)[:18000]


def set_ai_selection(label: str, text: str):
    st.session_state.selected_text_for_ai = f"{label}:\n{text or ''}"[:5000]


def format_quiz_question_for_context(question: dict, number: int) -> str:
    options = "\n".join(f"- {option}" for option in question.get("options", []))
    return (
        f"{number}. {question.get('question', '')}\n"
        f"Options:\n{options}\n"
        f"Correct answer: {question.get('correct_answer', '')}\n"
        f"Explanation: {question.get('explanation', '')}"
    )


def format_exam_for_context(tasks: list) -> str:
    lines = []
    for index, task in enumerate(tasks, 1):
        lines.append(f"{index}. {task.get('title', 'Task')} ({task.get('difficulty', '')}, {task.get('points', 0)} points)")
        lines.append(task.get("prompt", ""))
        lines.append(f"Expected answer: {task.get('expected_answer', '')}")
        lines.append("")
    return "\n".join(lines)


def extract_text_from_file(path: str, suffix: str) -> str:
    if suffix in [".txt", ".md"]:
        return Path(path).read_text(errors="ignore")
    if suffix == ".docx":
        from docx import Document
        document = Document(path)
        return "\n".join(paragraph.text for paragraph in document.paragraphs)
    if suffix == ".pdf":
        from pypdf import PdfReader
        reader = PdfReader(path)
        return "\n".join(page.extract_text() or "" for page in reader.pages)
    if suffix == ".pptx":
        from pptx import Presentation
        presentation = Presentation(path)
        chunks = []
        for slide_index, slide in enumerate(presentation.slides, 1):
            chunks.append(f"Slide {slide_index}")
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text:
                    chunks.append(shape.text)
            chunks.append("")
        return "\n".join(chunks)
    raise ValueError(f"Unsupported file type: {suffix}")


def extract_text_from_image(path: str) -> str:
    client = ollama.Client(host=config.OLLAMA_HOST)
    response = client.chat(
        model=config.OLLAMA_VISION_MODEL,
        messages=[
            {
                "role": "user",
                "content": (
                    "Extract all lecture content from this image. "
                    "Keep the original language. Include headings, formulas, lists, and visible notes. "
                    "Return only the extracted lecture text."
                ),
                "images": [path],
            }
        ],
        options={"temperature": 0.1, "num_predict": 1800},
    )
    return response["message"]["content"]


def format_quiz_for_download(questions: list) -> str:
    lines = ["# Practice Quiz", ""]
    for i, question in enumerate(questions, 1):
        lines.append(f"## {i}. {question.get('question', '')}")
        for option in question.get("options", []):
            lines.append(f"- {option}")
        lines.append(f"Correct answer: {question.get('correct_answer', '')}")
        lines.append(f"Explanation: {question.get('explanation', '')}")
        lines.append("")
    return "\n".join(lines)


def format_flashcards_for_download(cards: list) -> str:
    lines = ["# Study Flashcards", ""]
    for i, card in enumerate(cards, 1):
        lines.append(f"## Card {i}")
        lines.append(f"Topic: {card.get('topic', '')}")
        lines.append(f"Front: {card.get('front', '')}")
        lines.append(f"Back: {card.get('back', '')}")
        lines.append("")
    return "\n".join(lines)


if __name__ == "__main__":
    main()
